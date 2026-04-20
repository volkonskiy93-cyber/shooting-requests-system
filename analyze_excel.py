# -*- coding: utf-8 -*-
import os
import sys

try:
    import pandas as pd
    import openpyxl
except ImportError:
    print("Устанавливаю необходимые библиотеки...")
    os.system("pip install pandas openpyxl")
    import pandas as pd
    import openpyxl

# Поиск файла
root = r'c:\Исходники DaVinci'
excel_file = None

for r, d, files in os.walk(root):
    for f in files:
        if 'прошлое' in f.lower() and f.endswith(('.xlsx', '.xls', '.xlsm')):
            excel_file = os.path.join(r, f)
            break
    if excel_file:
        break

if not excel_file:
    # Попробуем найти в папке выработка
    for r, d, files in os.walk(root):
        if 'выработка' in r.lower():
            for f in files:
                if f.endswith(('.xlsx', '.xls', '.xlsm')):
                    excel_file = os.path.join(r, f)
                    print(f"Найден Excel файл в папке выработка: {excel_file}")
                    break
            if excel_file:
                break

if not excel_file:
    print("Excel файл не найден!")
    sys.exit(1)

print(f"Анализирую файл: {excel_file}\n")

# Читаем Excel файл
try:
    # Открываем файл с openpyxl для анализа структуры
    wb = openpyxl.load_workbook(excel_file, data_only=True)
    
    print("=" * 80)
    print("СТРУКТУРА ФАЙЛА")
    print("=" * 80)
    print(f"Листы в файле: {wb.sheetnames}\n")
    
    # Анализируем каждый лист
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        print(f"\n{'=' * 80}")
        print(f"ЛИСТ: {sheet_name}")
        print(f"{'=' * 80}")
        
        # Проверяем наличие автофильтров
        if sheet.auto_filter:
            print(f"✓ Автофильтр включен на диапазоне: {sheet.auto_filter.ref}")
        
        # Читаем данные с pandas для анализа
        try:
            df = pd.read_excel(excel_file, sheet_name=sheet_name, header=0)
            
            print(f"\nРазмер данных: {df.shape[0]} строк, {df.shape[1]} столбцов")
            print(f"\nСтолбцы:")
            for i, col in enumerate(df.columns, 1):
                print(f"  {i}. {col}")
            
            # Ищем столбцы с фильтрами (корреспонденты, редакторы и т.д.)
            print(f"\nАнализ столбцов на наличие фильтров:")
            for col in df.columns:
                col_lower = str(col).lower()
                if any(word in col_lower for word in ['корреспондент', 'редактор', 'фильтр', 'filter']):
                    unique_vals = df[col].dropna().unique()
                    print(f"  - {col}: {len(unique_vals)} уникальных значений")
                    if len(unique_vals) <= 20:
                        print(f"    Значения: {', '.join(map(str, unique_vals[:20]))}")
            
            # Показываем первые несколько строк
            print(f"\nПервые 5 строк данных:")
            print(df.head().to_string())
            
            # Проверяем наличие пустых значений
            print(f"\nПустые значения по столбцам:")
            null_counts = df.isnull().sum()
            for col, count in null_counts.items():
                if count > 0:
                    print(f"  {col}: {count} пустых значений")
        
        except Exception as e:
            print(f"Ошибка при чтении данных pandas: {e}")
            # Показываем данные через openpyxl
            print(f"\nДанные (первые 10 строк):")
            for row_idx, row in enumerate(sheet.iter_rows(min_row=1, max_row=11, values_only=True), 1):
                print(f"Строка {row_idx}: {row}")
    
    wb.close()
    
except Exception as e:
    print(f"Ошибка при анализе файла: {e}")
    import traceback
    traceback.print_exc()
