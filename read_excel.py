# -*- coding: utf-8 -*-
"""
Скрипт для анализа Excel файла "прошлое время" из папки "выработка"
Изучает структуру файла, фильтры по корреспондентам, редакторам и т.д.
"""
import os
import sys

try:
    import pandas as pd
    import openpyxl
except ImportError:
    print("Устанавливаю необходимые библиотеки...")
    os.system("pip install pandas openpyxl")
    import pandas as pd
    import openpyxl

def analyze_excel_file(file_path):
    """Анализирует Excel файл и выводит информацию о структуре и фильтрах"""
    
    if not os.path.exists(file_path):
        print(f"Файл не найден: {file_path}")
        return
    
    print("=" * 100)
    print(f"АНАЛИЗ ФАЙЛА: {os.path.basename(file_path)}")
    print("=" * 100)
    
    try:
        # Открываем файл с openpyxl для анализа структуры
        wb = openpyxl.load_workbook(file_path, data_only=True)
        
        print(f"\n📊 ЛИСТЫ В ФАЙЛЕ: {len(wb.sheetnames)}")
        for i, sheet_name in enumerate(wb.sheetnames, 1):
            print(f"   {i}. {sheet_name}")
        
        # Анализируем каждый лист
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            print(f"\n{'=' * 100}")
            print(f"📋 ЛИСТ: {sheet_name}")
            print(f"{'=' * 100}")
            
            # Проверяем наличие автофильтров
            if sheet.auto_filter:
                print(f"\n✅ АВТОФИЛЬТР ВКЛЮЧЕН")
                print(f"   Диапазон: {sheet.auto_filter.ref}")
                
                # Анализируем фильтры по столбцам
                if hasattr(sheet.auto_filter, 'filterColumn'):
                    print(f"   Столбцы с фильтрами: {len(sheet.auto_filter.filterColumn)}")
            
            # Читаем данные с pandas для анализа
            try:
                df = pd.read_excel(file_path, sheet_name=sheet_name, header=0)
                
                print(f"\n📈 РАЗМЕР ДАННЫХ:")
                print(f"   Строк: {df.shape[0]}")
                print(f"   Столбцов: {df.shape[1]}")
                
                print(f"\n📑 СТОЛБЦЫ:")
                for i, col in enumerate(df.columns, 1):
                    non_null = df[col].notna().sum()
                    print(f"   {i:2d}. {col:<40} (заполнено: {non_null}/{len(df)})")
                
                # Ищем столбцы с фильтрами (корреспонденты, редакторы и т.д.)
                print(f"\n🔍 АНАЛИЗ ФИЛЬТРОВ:")
                filter_columns = []
                for col in df.columns:
                    col_lower = str(col).lower()
                    if any(word in col_lower for word in ['корреспондент', 'редактор', 'фильтр', 'filter', 'автор', 'writer']):
                        filter_columns.append(col)
                        unique_vals = df[col].dropna().unique()
                        print(f"\n   📌 {col}:")
                        print(f"      Уникальных значений: {len(unique_vals)}")
                        if len(unique_vals) <= 30:
                            print(f"      Значения:")
                            for val in sorted(unique_vals)[:30]:
                                count = (df[col] == val).sum()
                                print(f"         - {val} ({count} записей)")
                        else:
                            print(f"      Первые 20 значений:")
                            for val in sorted(unique_vals)[:20]:
                                count = (df[col] == val).sum()
                                print(f"         - {val} ({count} записей)")
                            print(f"      ... и еще {len(unique_vals) - 20} значений")
                
                if not filter_columns:
                    print("   Фильтры по корреспондентам/редакторам не найдены в названиях столбцов")
                    print("   Проверяю все столбцы на наличие категориальных данных...")
                    for col in df.columns:
                        if df[col].dtype == 'object':
                            unique_count = df[col].nunique()
                            total_count = len(df[col].dropna())
                            if 2 <= unique_count <= 50 and total_count > 0:
                                print(f"\n   📌 {col} (возможный фильтр):")
                                print(f"      Уникальных значений: {unique_count}")
                                print(f"      Примеры: {', '.join(map(str, df[col].dropna().unique()[:10]))}")
                
                # Показываем первые несколько строк
                print(f"\n📄 ПЕРВЫЕ 5 СТРОК ДАННЫХ:")
                print(df.head().to_string())
                
                # Статистика по числовым столбцам
                numeric_cols = df.select_dtypes(include=['number']).columns
                if len(numeric_cols) > 0:
                    print(f"\n📊 СТАТИСТИКА ПО ЧИСЛОВЫМ СТОЛБЦАМ:")
                    print(df[numeric_cols].describe().to_string())
                
                # Проверяем наличие пустых значений
                null_counts = df.isnull().sum()
                if null_counts.sum() > 0:
                    print(f"\n⚠️  ПУСТЫЕ ЗНАЧЕНИЯ:")
                    for col, count in null_counts.items():
                        if count > 0:
                            percentage = (count / len(df)) * 100
                            print(f"   {col}: {count} ({percentage:.1f}%)")
            
            except Exception as e:
                print(f"⚠️  Ошибка при чтении данных pandas: {e}")
                import traceback
                traceback.print_exc()
                # Показываем данные через openpyxl
                print(f"\n📄 ДАННЫЕ (первые 10 строк через openpyxl):")
                for row_idx, row in enumerate(sheet.iter_rows(min_row=1, max_row=11, values_only=True), 1):
                    print(f"   Строка {row_idx}: {row}")
        
        wb.close()
        print(f"\n{'=' * 100}")
        print("✅ АНАЛИЗ ЗАВЕРШЕН")
        print(f"{'=' * 100}")
        
    except Exception as e:
        print(f"❌ Ошибка при анализе файла: {e}")
        import traceback
        traceback.print_exc()

if __name__ == "__main__":
    # Пробуем найти файл автоматически
    if len(sys.argv) > 1:
        file_path = sys.argv[1]
    else:
        # Автопоиск
        root = r'c:\Исходники DaVinci'
        file_path = None
        
        print("🔍 Поиск файла...")
        for r, d, files in os.walk(root):
            # Ищем файл с "прошлое" в названии
            for f in files:
                if 'прошлое' in f.lower() and f.endswith(('.xlsx', '.xls', '.xlsm')):
                    file_path = os.path.join(r, f)
                    print(f"✅ Найден файл: {file_path}")
                    break
            if file_path:
                break
            
            # Ищем в папке выработка
            if not file_path and 'выработка' in r.lower():
                for f in files:
                    if f.endswith(('.xlsx', '.xls', '.xlsm')):
                        file_path = os.path.join(r, f)
                        print(f"✅ Найден Excel файл в папке выработка: {file_path}")
                        break
                if file_path:
                    break
        
        if not file_path:
            print("❌ Файл не найден!")
            print("\nИспользование:")
            print(f"  python {sys.argv[0]} <путь_к_файлу>")
            print("\nПример:")
            print(f"  python {sys.argv[0]} \"c:\\Исходники DaVinci\\дут\\выработка\\прошлое время.xlsx\"")
            sys.exit(1)
    
    analyze_excel_file(file_path)
