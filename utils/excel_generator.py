"""
Генерация Excel файлов для заявок ФИГАРО и ТТК на основе шаблонов.
Особое внимание уделено сохранению стилей и правильным координатам.
"""

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Font, Color
from datetime import datetime, date
import os
import tempfile
import re
from pathlib import Path

def _parse_date_yyyy_mm_dd(date_string):
    if not date_string:
        return None
    if isinstance(date_string, (date, datetime)):
        return date_string
    try:
        return datetime.strptime(str(date_string), "%Y-%m-%d").date()
    except:
        return None

def _parse_datetime_local(dt_string):
    if not dt_string:
        return None
    try:
        return datetime.strptime(str(dt_string), "%Y-%m-%dT%H:%M")
    except:
        return _parse_date_yyyy_mm_dd(dt_string)

def _template_path(contractor: str, without_main_kit: bool = False) -> Path:
    project_root = Path(__file__).resolve().parent.parent
    if contractor == "ttk" and without_main_kit:
        return project_root / "excel_templates" / "ttk_without_main_kit_template.xlsx"
    return project_root / "excel_templates" / f"{contractor}_template.xlsx"

def _set_cell_value(ws, coord_or_cell, value, force_black=True):
    """
    Безопасно устанавливает значение ячейки, даже если она объединена.
    По возможности сохраняет шрифт или принудительно ставит черный цвет.
    """
    if value is None:
        return

    if isinstance(coord_or_cell, str):
        cell = ws[coord_or_cell]
    else:
        cell = coord_or_cell

    target_cell = cell
    if isinstance(cell, MergedCell):
        # Если ячейка объединена, ищем начало диапазона
        for merged_range in ws.merged_cells.ranges:
            if cell.coordinate in merged_range:
                target_cell = ws.cell(row=merged_range.min_row, column=merged_range.min_col)
                break
    
    target_cell.value = value
    
    # Исправление "красного шрифта": принудительно ставим черный цвет, если нужно
    if force_black and target_cell.font:
        new_font = Font(
            name=target_cell.font.name,
            size=target_cell.font.size,
            bold=target_cell.font.bold,
            italic=target_cell.font.italic,
            color="000000" # Чистый черный
        )
        target_cell.font = new_font

def _resolve_target_cell(ws, coord_or_cell):
    if isinstance(coord_or_cell, str):
        cell = ws[coord_or_cell]
    else:
        cell = coord_or_cell

    if isinstance(cell, MergedCell):
        for merged_range in ws.merged_cells.ranges:
            if cell.coordinate in merged_range:
                return ws.cell(row=merged_range.min_row, column=merged_range.min_col)
    return cell

def _clear_cell_value(ws, coord_or_cell):
    target_cell = _resolve_target_cell(ws, coord_or_cell)
    if target_cell is not None:
        target_cell.value = None


def _set_cell_date(ws, coord_or_cell, value, number_format="dd.mm.yy"):
    parsed_date = _parse_date_yyyy_mm_dd(value)
    if not parsed_date:
        return

    _set_cell_value(ws, coord_or_cell, parsed_date, force_black=False)
    target_cell = _resolve_target_cell(ws, coord_or_cell)
    if target_cell is not None:
        target_cell.number_format = number_format


def _format_extension_value(value):
    normalized = str(value or "").strip()
    return normalized or "нет"

def _replace_qty_in_text(text: str, qty: int) -> str:
    if not isinstance(text, str):
        return text
    try:
        qty_int = int(qty or 0)
        # Ищем паттерны типа -2шт, 2 шт, -2 шт.
        pat = re.compile(r"([- ]?)(\d+)(\s*шт)", re.IGNORECASE)
        if pat.search(text):
            return pat.sub(rf"\1{qty_int}\3", text, count=1)
        return text
    except:
        return text

def _apply_equipment(ws, equipment, contractor=None, without_main_kit=False):
    """
    Заполняет таблицу оборудования. 
    В шаблонах ФИГАРО/ТТК: 
    A - Основной комплект, C - Выбор (да/нет), D - Количество доп.
    """
    if not equipment:
        return
    
    # Ищем строку заголовка оборудования
    header_row = None
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if isinstance(a, str) and (
            "Основной комплект" in a or "Без основного комплекта" in a
        ):
            header_row = r
            break
    
    if not header_row:
        return

    start_row = header_row + 1

    if contractor == 'ttk' and without_main_kit:
        for row_idx in range(start_row, ws.max_row + 1):
            cell_main = ws.cell(row=row_idx, column=1)
            cell_add = ws.cell(row=row_idx, column=2)
            cell_choice = ws.cell(row=row_idx, column=3)
            cell_qty = ws.cell(row=row_idx, column=4)

            main_text = cell_main.value if isinstance(cell_main.value, str) else ''
            add_text = cell_add.value if isinstance(cell_add.value, str) else ''

            if not main_text and not add_text:
                break

            if main_text and not add_text:
                _set_cell_value(ws, cell_add, "—", force_black=False)
                _clear_cell_value(ws, cell_choice)
                _clear_cell_value(ws, cell_qty)

    for idx, eq in enumerate(equipment):
        # Если есть rowIndex (из UI), используем его, иначе используем порядковый номер
        ui_row_index = eq.get("rowIndex")
        if ui_row_index is not None:
            row_idx = start_row + ui_row_index
        else:
            row_idx = start_row + idx
            
        if row_idx > ws.max_row:
            continue

        main_qty = eq.get("mainQuantity", 0)
        add_qty = eq.get("additionalQuantity", 0)

        # Колонка A: Обновляем текст основного комплекта.
        # Для ТТК в режиме "без основного комплекта" не затираем шаблон:
        # в образце основной блок остается как в исходном файле.
        cell_main = ws.cell(row=row_idx, column=1)
        if cell_main.value and isinstance(cell_main.value, str):
            if not (contractor == 'ttk' and without_main_kit):
                new_val = _replace_qty_in_text(cell_main.value, main_qty)
                _set_cell_value(ws, cell_main, new_val, force_black=False) # Сохраняем стиль шаблона

        # Колонка C: да/нет
        cell_choice = ws.cell(row=row_idx, column=3)
        has_additional = int(add_qty or 0) > 0
        if contractor == 'ttk':
            if without_main_kit:
                _set_cell_value(ws, cell_choice, "ДА" if has_additional else "НЕТ", force_black=False)
            else:
                _set_cell_value(ws, cell_choice, "ДА" if has_additional else "НЕТ", force_black=False)
        elif cell_choice.value is not None:
            _set_cell_value(ws, cell_choice, "да" if has_additional else "нет", force_black=False)

        # Колонка D: количество
        cell_qty = ws.cell(row=row_idx, column=4)
        if int(add_qty or 0) > 0:
            _set_cell_value(ws, cell_qty, int(add_qty), force_black=False)
        else:
            # Если 0, лучше очистить ячейку количества, чтобы не было лишних нулей
            if cell_qty.value is not None and not isinstance(cell_qty, MergedCell):
                cell_qty.value = None

def create_excel_document(form_data, application_id):
    contractor = form_data.get('contractor')
    without_main_kit = bool(form_data.get("withoutMainKit"))
    tpl = _template_path(contractor, without_main_kit=without_main_kit)
    
    if not tpl.exists():
        print(f"⚠️ Шаблон не найден: {tpl}")
        return None

    try:
        # data_only=False чтобы сохранить формулы и стили
        wb = openpyxl.load_workbook(tpl)
        ws = wb.active

        if contractor == 'figaro':
            # Точные координаты для ФИГАРО
            _set_cell_value(ws, "B2", form_data.get("storyTitle", ""))
            _set_cell_value(ws, "B3", form_data.get("annotation", ""))
            _set_cell_value(ws, "B4", form_data.get("notes", ""))
            _set_cell_value(ws, "B5", "да" if form_data.get("accreditation") == "yes" else "нет")
            _set_cell_value(ws, "B7", form_data.get("director", ""))
            _set_cell_value(ws, "B8", form_data.get("correspondent", ""))
            _set_cell_value(ws, "B9", form_data.get("producer", ""))
            
            op = form_data.get("operator", "")
            ve = form_data.get("videoEngineer", "")
            _set_cell_value(ws, "B10", f"{op}, {ve}".strip(", "))
            
            _set_cell_date(ws, "A13", form_data.get("applicationDate"))
            _set_cell_date(ws, "B13", form_data.get("shootingDate"))
            _set_cell_value(ws, "C13", f"{form_data.get('startTime')} - {form_data.get('endTime')}")
            _set_cell_value(ws, "D13", _parse_date_yyyy_mm_dd(form_data.get("broadcastDate")))
            
            _apply_equipment(ws, form_data.get("equipment", []), contractor='figaro')

        elif contractor == 'ttk':
            # Точные координаты для ТТК (на основе дампа)
            _set_cell_value(ws, "B2", form_data.get("storyTitle", ""))
            _set_cell_value(ws, "B3", form_data.get("annotation", ""))
            _set_cell_value(ws, "B4", form_data.get("notes", ""))
            _set_cell_value(ws, "B5", "да" if form_data.get("accreditation") == "yes" else "нет")
            
            # Ответственные ТТК
            _set_cell_value(ws, "B7", form_data.get("director", ""))
            _set_cell_value(ws, "B8", form_data.get("correspondent", ""))
            _set_cell_value(ws, "B9", form_data.get("producer", ""))
            
            op = form_data.get("operator", "")
            ve = form_data.get("videoEngineer", "")
            _set_cell_value(ws, "B10", f"{op}, {ve}".strip(", "))
            
            # Даты и время (строка 14 для значений)
            _set_cell_date(ws, "A14", form_data.get("applicationDate"))
            _set_cell_date(ws, "B14", form_data.get("shootingDate"))
            _set_cell_value(ws, "C14", f"{form_data.get('startTime')} - {form_data.get('endTime')}")
            _set_cell_value(ws, "D14", _parse_date_yyyy_mm_dd(form_data.get("broadcastDate")))
            _set_cell_value(ws, "D15", _format_extension_value(form_data.get("extension", "")))
            
            # Доп поля ТТК
            # В ТТК "Уточнения" обычно в C5 (значение в D5 или B5?)
            # Но по дампу C5 - это заголовок. Пробуем писать в D5.
            if form_data.get("clarifications"):
                _set_cell_value(ws, "D5", form_data.get("clarifications"))
            
            # Сдача материала (ТТК) - по дампу это может быть ниже
            # Если не знаем точно, ищем по метке, но с вертикальным смещением +1
            def _set_below_label(label, val):
                for row in ws.iter_rows():
                    for cell in row:
                        if isinstance(cell.value, str) and label.lower() in cell.value.lower():
                            _set_cell_value(ws, ws.cell(row=cell.row + 1, column=cell.column), val)
                            return True
                return False

            _apply_equipment(
                ws,
                form_data.get("equipment", []),
                contractor='ttk',
                without_main_kit=without_main_kit
            )

        # Сохранение
        temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
        wb.save(temp_file.name)
        temp_file.close()
        return temp_file.name

    except Exception as e:
        print(f"❌ Ошибка генерации Excel: {e}")
        import traceback
        traceback.print_exc()
        return None
