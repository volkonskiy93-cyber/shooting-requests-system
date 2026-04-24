"""
Генерация Word документов для заявок продюсерам
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import datetime
import tempfile


def format_date(date_string):
    """Форматирование даты из строки"""
    if not date_string:
        return ''
    try:
        date = datetime.strptime(date_string, '%Y-%m-%d')
        return date.strftime('%d.%m.%Y')
    except:
        return date_string


def format_broadcast_date(date_string):
    """Дата эфира или значение по умолчанию."""
    return format_date(date_string) or 'по гот.'


def create_word_document(form_data, application_id):
    """
    Создание Word документа для заявки продюсерам
    
    Args:
        form_data: словарь с данными формы
        application_id: ID заявки
    
    Returns:
        путь к созданному файлу или None
    """
    if form_data.get('contractor') != 'producer':
        return None
    
    doc = Document()
    
    # Настройка стилей
    title_style = doc.styles['Heading 1']
    title_font = title_style.font
    title_font.name = 'Arial'
    title_font.size = Pt(16)
    title_font.bold = True
    
    # Заголовок
    title = doc.add_heading('Заявка продюсерам', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    subtitle = doc.add_paragraph('Программа "Доброе утро"')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_format = subtitle.runs[0].font
    subtitle_format.name = 'Arial'
    subtitle_format.size = Pt(12)
    
    doc.add_paragraph()  # Пустая строка
    
    # Основная информация
    fields = [
        ('Название сюжета', form_data.get('storyTitle', '')),
        ('Дата эфира', format_broadcast_date(form_data.get('broadcastDate'))),
        ('Краткое содержание сюжета', form_data.get('summary', '')),
        ('Разработка от отдела планирования', form_data.get('planningDevelopment', '')),
        ('Герои', form_data.get('heroes', '')),
        ('Дата съемки', format_date(form_data.get('shootingDate'))),
        ('Корреспондент', form_data.get('correspondentText') or form_data.get('correspondent', '')),
        ('Контакты корреспондента', form_data.get('correspondentContacts', '')),
        ('Редактор', form_data.get('directorText') or form_data.get('director', '')),
        ('Дата подачи заявки', format_date(form_data.get('applicationDate'))),
    ]
    
    for label, value in fields:
        p = doc.add_paragraph()
        run_label = p.add_run(f'{label}: ')
        run_label.font.name = 'Arial'
        run_label.font.size = Pt(11)
        run_label.bold = True
        
        run_value = p.add_run(value or '-')
        run_value.font.name = 'Arial'
        run_value.font.size = Pt(11)
    
    # Номер заявки
    doc.add_paragraph()
    p = doc.add_paragraph()
    run = p.add_run(f'Номер заявки: {str(application_id)[:8]}')
    run.font.name = 'Arial'
    run.font.size = Pt(10)
    run.font.italic = True
    
    # Сохранение во временный файл
    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix='.docx')
    doc.save(temp_file.name)
    temp_file.close()
    
    return temp_file.name
